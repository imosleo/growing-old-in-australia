"""Parse an ABS time-series workbook (Index + Data sheets) into long format.

ABS time-series workbooks share a fixed shape:
  Index sheet  : 'Data Item Description' + 'Series ID' from row 11 down
  Data* sheets : row 10 holds Series IDs, rows 11+ hold date in col A then values
"""
import re
import sys
import openpyxl
import pandas as pd


def load_abs(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    # Series ID -> full description
    desc = {}
    ws = wb["Index"]
    for row in ws.iter_rows(min_row=11, max_col=8, values_only=True):
        label, sid = row[0], row[4]
        if label and sid:
            desc[str(sid).strip()] = str(label).strip()

    records = []
    for name in wb.sheetnames:
        if not name.startswith("Data"):
            continue
        ws = wb[name]
        rows = list(ws.iter_rows(values_only=True))
        header = None
        for r in rows:
            if r and str(r[0]).strip() == "Series ID":
                header = r
                break
        if header is None:
            continue
        for r in rows:
            if not r or not hasattr(r[0], "year"):
                continue
            year = r[0].year
            for cix in range(1, len(header)):
                sid = header[cix]
                if not sid or cix >= len(r):
                    continue
                val = r[cix]
                if val is None or val == "":
                    continue
                records.append((str(sid).strip(), year, val))

    wb.close()
    df = pd.DataFrame(records, columns=["series_id", "year", "value"])
    df["description"] = df["series_id"].map(desc)
    return df


if __name__ == "__main__":
    path = sys.argv[1]
    df = load_abs(path)
    print("rows:", len(df))
    print("years:", df["year"].min(), "-", df["year"].max())
    print("series:", df["series_id"].nunique())
    print("\nsample descriptions:")
    for d in pd.Series(df["description"].dropna().unique()).head(12):
        print("   ", d)
    print("\nunmapped series ids:", df["description"].isna().sum())
    # show the shape of the description string
    uniq = sorted(set(re.sub(r"\d+", "#", str(d)) for d in df["description"].dropna().unique()))
    print("\ndescription patterns (%d):" % len(uniq))
    for u in uniq[:20]:
        print("   ", u)
