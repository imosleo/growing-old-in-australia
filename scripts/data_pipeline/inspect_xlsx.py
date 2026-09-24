"""List sheets and preview the top-left of each sheet in an xlsx workbook."""
import sys
import openpyxl

path = sys.argv[1]
maxrows = int(sys.argv[2]) if len(sys.argv) > 2 else 8
maxcols = int(sys.argv[3]) if len(sys.argv) > 3 else 8

wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
print("WORKBOOK:", path)
print("SHEETS:", len(wb.sheetnames))
for name in wb.sheetnames:
    ws = wb[name]
    print("\n--- [%s]  dims=%s ---" % (name, ws.calculate_dimension()))
    for rix, row in enumerate(ws.iter_rows(max_row=maxrows, max_col=maxcols, values_only=True)):
        cells = ["" if v is None else str(v)[:38] for v in row]
        if any(cells):
            print("  r%-2d | %s" % (rix + 1, " | ".join(cells)))
wb.close()
